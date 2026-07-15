import os
from pathlib import Path
import docx
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

def create_docx_resume(dest_path: Path):
    """Creates a sample Word (.docx) resume for John Doe."""
    doc = docx.Document()
    
    # Title
    doc.add_heading("John Doe", level=0)
    
    # Contact Info
    p = doc.add_paragraph()
    p.add_run("Email: john.doe@email.com | Phone: +1-555-0199 | Address: New York, NY\n")
    p.add_run("LinkedIn: linkedin.com/in/johndoe | GitHub: github.com/johndoe")
    
    # Summary
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph(
        "Results-driven Software Engineer with 3 years of professional experience building scalable web applications and machine learning integrations. "
        "Proficient in Python, Django, React, and MySQL, with a solid understanding of software design patterns and cloud services."
    )
    
    # Experience
    doc.add_heading("Work Experience", level=1)
    
    p_job1 = doc.add_paragraph()
    p_job1.add_run("Software Engineer | WebTech Solutions\n").bold = True
    p_job1.add_run("June 2024 - Present | New York, NY\n").italic = True
    p_job1.add_run("• Developed and maintained backend services using Python, Django, and PostgreSQL.\n")
    p_job1.add_run("• Designed and implemented responsive UI components in React, improving user engagement by 20%.\n")
    p_job1.add_run("• Optimized database queries, reducing API response times by 35%.")
    
    p_job2 = doc.add_paragraph()
    p_job2.add_run("Junior Developer | CodeCraft Inc.\n").bold = True
    p_job2.add_run("August 2023 - May 2024 | Boston, MA\n").italic = True
    p_job2.add_run("• Collaborated with a team of 5 developers to build and test features for a customer portal using Python and Flask.\n")
    p_job2.add_run("• Implemented RESTful APIs and integrated third-party payment gateways.\n")
    p_job2.add_run("• Authored unit tests to achieve 85% code coverage.")
    
    # Skills
    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph(
        "Languages: Python, JavaScript, HTML5, CSS3, SQL\n"
        "Frameworks & Tools: Django, Flask, React, Node.js, Git, Docker\n"
        "Databases: MySQL, PostgreSQL, SQLite"
    )
    
    # Projects
    doc.add_heading("Key Projects", level=1)
    
    p_proj1 = doc.add_paragraph()
    p_proj1.add_run("Stock Trend Predictor (Machine Learning Project)\n").bold = True
    p_proj1.add_run("• Developed a Python-based machine learning pipeline using pandas, scikit-learn, and Flask to predict stock price movements.\n")
    p_proj1.add_run("• Trained an LSTM neural network model on historical stock price data, achieving a 78% prediction accuracy.\n")
    p_proj1.add_run("• Created an interactive web interface for users to visualize predictions.")
    
    p_proj2 = doc.add_paragraph()
    p_proj2.add_run("SaaS Analytics Dashboard\n").bold = True
    p_proj2.add_run("• Built a real-time metrics dashboard using React, TailwindCSS, Node.js, and MongoDB.\n")
    p_proj2.add_run("• Integrated WebSocket for live chart updates and implemented JWT-based user authentication.")
    
    # Education
    doc.add_heading("Education", level=1)
    p_edu = doc.add_paragraph()
    p_edu.add_run("B.S. in Computer Science | Boston University\n").bold = True
    p_edu.add_run("Graduated: May 2023").italic = True
    
    doc.save(dest_path)
    print(f"Created DOCX resume at: {dest_path}")

def create_pdf_resume(dest_path: Path):
    """Creates a sample PDF resume for Sarah Smith using ReportLab."""
    doc = SimpleDocTemplate(str(dest_path), pagesize=letter,
                            rightMargin=54, leftMargin=54,
                            topMargin=54, bottomMargin=54)
    
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'NameHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=24,
        leading=28,
        textColor=colors.HexColor('#2c3e50'),
        spaceAfter=6
    )
    contact_style = ParagraphStyle(
        'ContactInfo',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#7f8c8d'),
        spaceAfter=15
    )
    h1_style = ParagraphStyle(
        'SectionHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#2980b9'),
        spaceBefore=12,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        'BodyTextCustom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        spaceAfter=6
    )
    bullet_style = ParagraphStyle(
        'BulletCustom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        leftIndent=15,
        firstLineIndent=-10,
        spaceAfter=4
    )
    
    story = []
    
    # Name and title
    story.append(Paragraph("Sarah Smith", title_style))
    story.append(Paragraph("Data Scientist & ML Engineer", ParagraphStyle('SubTitle', fontName='Helvetica-Bold', fontSize=12, leading=14, textColor=colors.HexColor('#34495e'), spaceAfter=4)))
    story.append(Paragraph("Email: sarah.smith@email.com | Phone: +1-555-0255 | Boston, MA | github.com/sarahsmith", contact_style))
    
    # Summary
    story.append(Paragraph("Professional Summary", h1_style))
    story.append(Paragraph(
        "Highly skilled Data Scientist with 5 years of experience designing and deploying machine learning models, natural language processing pipelines, and deep learning neural networks. "
        "Expertise in Python, PyTorch, SQL, and cloud deployments (AWS). Passionate about turning complex data into actionable business intelligence.",
        body_style
    ))
    
    # Experience
    story.append(Paragraph("Professional Experience", h1_style))
    
    story.append(Paragraph("<b>Senior Data Scientist</b> | Analytics Corp (2023 - Present)", body_style))
    story.append(Paragraph("• Spearheaded the deployment of a transformer-based NLP chatbot that resolved 40% of customer support queries automatically.", bullet_style))
    story.append(Paragraph("• Developed a real-time anomaly detection system in Python and PyTorch, saving the company $150k annually in credit card fraud.", bullet_style))
    story.append(Paragraph("• Managed a team of 3 data scientists and established ML engineering best practices.", bullet_style))
    
    story.append(Spacer(1, 4))
    story.append(Paragraph("<b>Data Analyst</b> | Insight Data Partners (2021 - 2023)", body_style))
    story.append(Paragraph("• Built predictive models using scikit-learn to forecast customer churn with 85% recall.", bullet_style))
    story.append(Paragraph("• Extracted and processed large scale datasets using SQL and Spark pipelines.", bullet_style))
    
    # Skills
    story.append(Paragraph("Technical Skills", h1_style))
    story.append(Paragraph("<b>Languages:</b> Python, R, SQL, C++, Java", body_style))
    story.append(Paragraph("<b>Machine Learning:</b> Deep Learning, PyTorch, TensorFlow, NLP, Scikit-Learn, Pandas, NumPy", body_style))
    story.append(Paragraph("<b>Cloud & Databases:</b> AWS (S3, SageMaker, EC2), PostgreSQL, MongoDB, Git, Docker", body_style))
    
    # Projects
    story.append(Paragraph("Projects", h1_style))
    story.append(Paragraph("<b>AI-Driven Fraud Detection Engine</b>", body_style))
    story.append(Paragraph("• Designed a deep learning autoencoder neural network in PyTorch to detect fraudulent transactions.", bullet_style))
    story.append(Paragraph("• Processed a highly imbalanced dataset of 10M+ transactions, achieving an Area Under ROC curve of 0.96.", bullet_style))
    
    story.append(Spacer(1, 4))
    story.append(Paragraph("<b>Semantic Search Engine (NLP Project)</b>", body_style))
    story.append(Paragraph("• Developed a semantic search tool using SentenceTransformers and FAISS (Facebook AI Similarity Search) to retrieve documents based on user intent.", bullet_style))
    story.append(Paragraph("• Deployed the final engine as a containerized microservice on AWS ECS using FastAPI.", bullet_style))
    
    # Education
    story.append(Paragraph("Education", h1_style))
    story.append(Paragraph("<b>M.S. in Data Science</b> | Northeastern University (2021)", body_style))
    story.append(Paragraph("<b>B.S. in Mathematics</b> | University of Massachusetts (2019)", body_style))
    
    doc.build(story)
    print(f"Created PDF resume at: {dest_path}")

def generate_jds(samples_dir: Path):
    """Generates two sample job description files."""
    # ML JD
    jd_ml = (
        "Job Title: Machine Learning Engineer\n"
        "Experience Required: 3+ years of experience in data science or ML engineering\n"
        "Key Responsibilities:\n"
        "- Design, build, and deploy machine learning and deep learning models.\n"
        "- Develop natural language processing pipelines and predictive analytics applications.\n"
        "- Work with PyTorch, TensorFlow, and Python to train neural networks.\n"
        "- Query datasets using SQL and construct data pipelines.\n"
        "Required Skills:\n"
        "- Python\n"
        "- PyTorch or TensorFlow\n"
        "- SQL\n"
        "- Data Science (Scikit-Learn, Pandas)\n"
        "Project Requirements:\n"
        "- Must have built and deployed Machine Learning projects, such as fraud detection, NLP chatbot, or recommendation systems."
    )
    with open(samples_dir / "job_desc_ml.txt", "w", encoding="utf-8") as f:
        f.write(jd_ml)
        
    # Web JD
    jd_web = (
        "Job Title: Fullstack Web Developer (Python/React)\n"
        "Experience Required: 2+ years of software development experience\n"
        "Key Responsibilities:\n"
        "- Build responsive web interfaces and SaaS dashboards.\n"
        "- Develop backend API services in Python (Django or Flask).\n"
        "- Design database schemas in MySQL or PostgreSQL.\n"
        "Required Skills:\n"
        "- Python\n"
        "- React\n"
        "- JavaScript\n"
        "- Django\n"
        "- SQL (MySQL or PostgreSQL)\n"
        "Project Requirements:\n"
        "- Completed fullstack web applications, SaaS dashboards, or database-driven projects."
    )
    with open(samples_dir / "job_desc_web.txt", "w", encoding="utf-8") as f:
        f.write(jd_web)
        
    print(f"Created sample Job Descriptions in: {samples_dir}")

def main():
    samples_dir = Path(__file__).resolve().parent / "samples"
    samples_dir.mkdir(parents=True, exist_ok=True)
    
    create_docx_resume(samples_dir / "resume_john_doe.docx")
    create_pdf_resume(samples_dir / "resume_sarah_smith.pdf")
    generate_jds(samples_dir)
    print("All sample files generated successfully!")

if __name__ == "__main__":
    main()
